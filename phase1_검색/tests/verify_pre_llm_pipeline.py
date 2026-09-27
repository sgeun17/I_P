"""Latest file parser -> chunks -> real retrieval -> prompt -> fixture validation.

Run using phase1_검색/.venv Python from Search, passing --parser-python for a Python
with python-docx, openpyxl, python-pptx, reportlab and pdfplumber. No LLM or DB
upload service is called. PNG/JPG rows simulate OCR blocks, not image recognition.
"""
import argparse
from copy import deepcopy
from datetime import datetime
import hashlib
import json
from pathlib import Path
import subprocess
import sys

sys.dont_write_bytecode = True
ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT.parent / 'phase1_입력'
JUDGMENT = ROOT.parent / 'phase1_판단'


def save(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')


def fingerprint():
    paths = [ROOT/'controls.json']
    for folder in [INPUT, JUDGMENT, ROOT]:
        paths += [p for p in folder.rglob('*.py') if not any(x in p.parts for x in ('.venv','models','__pycache__'))]
    return {str(p.relative_to(ROOT.parent)): hashlib.sha256(p.read_bytes()).hexdigest() for p in paths}


def worker(out):
    sys.path.insert(0, str(INPUT))
    from parser_dispatcher import parse_file
    from chunking.chunker import make_chunks
    from chunking.chunk_format import validate_chunks
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from reportlab.pdfgen import canvas
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    fixtures = out/'fixtures'
    fixtures.mkdir(parents=True, exist_ok=True)
    text = '사용자 계정 관리 기록. 신규 계정은 업무별 최소권한을 확인하고 부서장의 승인 후 발급했다. 퇴직자 계정은 퇴직 당일 삭제했다. 계정별 신청자, 승인자, 처리일시와 권한 변경 내역을 기록했다.'
    specs = []
    for cid,name,content in [('TXT','account.txt',text.encode('utf-8')),
                              ('CP949','cp949.txt',text.encode('cp949')),
                              ('LONG','long.txt',((text+' ')*15).encode('utf-8')),
                              ('CSV','account.csv',('항목,내용\n계정관리,'+text+'\n').encode('utf-8-sig')),
                              ('UNRELATED','menu.txt','오늘 점심은 김치볶음밥과 달걀국이다.'.encode('utf-8')),
                              ('EMPTY','empty.txt',b''),('CORRUPTED','corrupted.txt',b'\x00\xff\x00binary')]:
        (fixtures/name).write_bytes(content)
        specs.append((cid,name))
    doc=Document()
    doc.add_heading('계정 관리',level=1)
    doc.add_paragraph(text)
    table=doc.add_table(rows=2,cols=2)
    for row,values in zip(table.rows,[['계정','처리'],['합성계정','퇴직 당일 삭제']]):
        for cell,value in zip(row.cells,values): cell.text=value
    doc.save(fixtures/'account.docx')
    specs.append(('DOCX','account.docx'))
    book=Workbook()
    for i,sheet in enumerate([book.active,book.create_sheet('승인 이력')],1):
        sheet.append(['항목','내용']); sheet.append([f'계정 {i}',text])
    book.save(fixtures/'account.xlsx')
    specs.append(('XLSX','account.xlsx'))
    deck=Presentation()
    for i in range(2):
        slide=deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text=f'계정 관리 {i+1}'
        slide.placeholders[1].text=text
    deck.save(fixtures/'account.pptx')
    specs.append(('PPTX','account.pptx'))
    # A local font keeps the PDF fixture self-contained and readable in Korean.
    font=Path('C:/Windows/Fonts/malgun.ttf')
    if font.is_file():
        pdfmetrics.registerFont(TTFont('FixtureKorean',str(font)))
    pdf=canvas.Canvas(str(fixtures/'account.pdf'))
    for n in (1,2):
        pdf.setFont('FixtureKorean' if font.is_file() else 'Helvetica',11)
        lines=[text[i:i+40] for i in range(0,len(text),40)] if font.is_file() else ['User accounts are issued after approval.', 'Permissions are revoked on employee departure.']
        for i,line in enumerate(lines): pdf.drawString(40,760-i*22,line)
        pdf.showPage()
    pdf.save()
    specs.append(('PDF','account.pdf'))
    rows=[]
    for n,(cid,name) in enumerate(specs,1):
        p=parse_file(fixtures/name)
        eid=f'E{8000+n:04d}'
        chunks=make_chunks(p,eid,2)
        assert not validate_chunks(chunks), (cid,validate_chunks(chunks))
        rows.append({'case_id':cid,'origin':'SYNTHETIC_FILE_ACTUAL_PARSER',
                     'fixture_sha256':hashlib.sha256((fixtures/name).read_bytes()).hexdigest(),
                     'payload':{'evidence_id':eid,'version':2,'source_file':name,
                                'file_type':p['file_type'],'chunks':chunks,'errors':p['errors']}})
    for n,kind in enumerate(['png','jpg'],21):
        p={'source_file':f'simulated.{kind}','file_type':kind,'errors':[],
           'blocks':[{'order':1,'block_type':'paragraph','level':None,'page':None,
                      'text':text,'table':None,'source':'ocr','confidence':.7}]}
        eid=f'E{8000+n:04d}'
        chunks=make_chunks(p,eid,2)
        assert not validate_chunks(chunks)
        assert all(c['source']=='ocr' for c in chunks)
        rows.append({'case_id':kind.upper()+'_CONTRACT','origin':'SIMULATED_OCR_BLOCKS_NO_IMAGE_RECOGNITION',
                     'payload':{'evidence_id':eid,'version':2,'source_file':p['source_file'],
                                'file_type':kind,'chunks':chunks,'errors':[]}})
    return rows


def main(parser_python,out):
    assert Path.cwd().resolve()==ROOT
    sys.path[:0]=[str(ROOT),str(JUDGMENT/'src'),str(ROOT/'tests')]
    from jsonschema import Draft202012Validator
    from chunk_retriever import LocalDocumentRetriever, prepare_input, RetrievalError, error_result
    from judgment_adapter import to_mapping_input, MappingAdapterError, read
    from retrieval_adapter import mapping_input_from_retriever
    from models import MappingInput
    from prompts import build_prompt_package
    from versions import build_version_info
    from service import build_result,to_evidence_status
    from enums import ErrorCode,ReviewReason
    from verify_judgment_integration import response_fixture
    before=fingerprint()
    out.mkdir(parents=True,exist_ok=True)
    save(out/'summary.json',{'status':'RUNNING','llm_calls':0})
    proc=subprocess.run([str(Path(parser_python).resolve()),'-B','-X','utf8',str(Path(__file__).resolve()),
                         '--worker','--output-dir',str(out)],capture_output=True,text=True,encoding='utf-8',timeout=120)
    if proc.returncode:
        raise RuntimeError(proc.stderr)
    records=json.loads(proc.stdout)
    save(out/'parsed_and_chunks.json',records)
    # Guard all invalid inputs before loading the model or issuing any search.
    rejected=[]
    valid=[]
    for record in records:
        cid,p=record['case_id'],record['payload']
        try:
            prepare_input(p)
            assert cid not in ('EMPTY','CORRUPTED')
            valid.append(record)
        except RetrievalError as e:
            assert cid in ('EMPTY','CORRUPTED') and e.code=='PREPROCESSING_FAILED'
            failure=error_result(e.code,str(e))
            try: to_mapping_input(failure)
            except MappingAdapterError: pass
            else: raise AssertionError('preprocessing failure reached judgment')
            rejected.append(cid)
            save(out/(cid+'_error.json'),failure)
    print('Parser/chunk checks complete; loading local retrieval model.',flush=True)
    engine=LocalDocumentRetriever()
    validator=Draft202012Validator(read(ROOT/'schemas/chunk_output.schema.json'))
    models,outputs,results={}, {}, []
    for record in valid:
        cid,p=record['case_id'],record['payload']
        result=engine.search(p,compatibility=False)
        validator.validate(result)
        assert result['warnings']==[] and result['evidence_chunks']==p['chunks']
        model=MappingInput.model_validate(to_mapping_input(result))
        assert mapping_input_from_retriever(result).model_dump()==model.model_dump()
        assert model.evidence_id==p['evidence_id'] and model.version==2
        assert [c.model_dump(mode='json') for c in model.chunks]==p['chunks']
        package=build_prompt_package(model)
        assert '"similarity_score"' not in package.user and '"rank"' not in package.user
        # Compare structured embedded evidence, avoiding mere substring tests.
        evidence=json.loads(package.user.split('<evidence>\n',1)[1].split('\n</evidence>',1)[0])
        assert [c['text'] for c in evidence['chunks']]==[c['text'] for c in p['chunks']]
        assert [c['source'] for c in evidence['chunks']]==[c['source'] for c in p['chunks']]
        save(out/(cid+'_search.json'),result)
        save(out/(cid+'_mapping_input.json'),model.model_dump(mode='json'))
        save(out/(cid+'_prompt.json'),{'mode':'PRE_LLM_NOT_SENT','messages':package.openai_compatible_messages(),
                                      'output_schema':package.output_schema,'prompt_version':package.prompt_version})
        models[cid],outputs[cid]=model,result
        row={'case_id':cid,'origin':record['origin'],'chunks':len(model.chunks),
             'pages':sorted({c.page_start for c in model.chunks if c.page_start is not None}),
             'sources':sorted({c.source.value for c in model.chunks}),
             'top5':[c.control_id for c in model.candidate_controls],'status':'PASS'}
        results.append(row)
        print(json.dumps(row,ensure_ascii=False),flush=True)
    assert [
        (c['control_id'],c['similarity_score']) for c in outputs['TXT']['retrieval']['candidates']]==[
        (c['control_id'],c['similarity_score']) for c in outputs['CP949']['retrieval']['candidates']]
    assert len(models['LONG'].chunks)>1
    for cid in ['PDF','XLSX','PPTX']:
        assert {p for c in models[cid].chunks for p in range(c.page_start,c.page_end+1)}=={1,2}, cid

    def response(cid,selected):
        return response_fixture(models[cid].model_dump(mode='json'),selected)

    chosen=[c.control_id for c in models['TXT'].candidate_controls]
    single=response('TXT',chosen[:1])
    bad=deepcopy(single)
    for branch in ('mapped_controls','candidate_decisions'):
        bad[branch][0]['citations'][0]['quote']='원문에 없는 허위 인용 PRELLM-INVALID'
    low=deepcopy(single)
    for branch in ('mapped_controls','candidate_decisions'): low[branch][0]['llm_confidence']=.4
    uncertain=response('TXT',[])
    uncertain['candidate_decisions'][0].update(decision='UNCERTAIN',llm_confidence=.4)
    scenarios=[('single','TXT',single,None),('multiple','TXT',response('TXT',chosen[:2]),None),
               ('no_match','UNRELATED',response('UNRELATED',[]),None),('low_confidence','TXT',low,None),
               ('uncertain','TXT',uncertain,None),('invalid_quote','TXT',bad,None),
               ('invalid_json','TXT','{broken',None),('call_failure','TXT',None,ErrorCode.LLM_RETRY_EXHAUSTED)]
    for cid in ['PNG_CONTRACT','JPG_CONTRACT']:
        scenarios.append((cid,cid,response(cid,[models[cid].candidate_controls[0].control_id]),None))
    checks=[]
    final_schema=Draft202012Validator(read(JUDGMENT/'schemas/phase1_mapping_result.schema.json'))
    for name,cid,fixture,error in scenarios:
        model=models[cid]
        raw=json.dumps(fixture,ensure_ascii=False) if isinstance(fixture,dict) else fixture
        r=build_result(raw,model,build_version_info(model,model_name='NO_LLM_FIXED_RESPONSE'),call_error=error)
        data=r.model_dump(mode='json')
        final_schema.validate(data)
        assert data['evidence_id']==model.evidence_id and data['version']==2
        if name in ('invalid_json','call_failure'):
            assert data['processing_status']=='FAILED' and r.human_review.required
        elif name=='invalid_quote':
            assert not r.validation.citations_valid and r.human_review.required
            assert any(i.code==ErrorCode.CITATION_QUOTE_NOT_IN_SOURCE for i in r.validation.issues)
        else:
            assert r.validation.passed and data['processing_status']=='COMPLETED'
            expected={'single':1,'multiple':2,'no_match':0,'uncertain':0,'low_confidence':1}.get(name,1)
            assert len(r.mapped_controls)==expected
        if name!='single': assert r.human_review.required, name
        if name.endswith('_CONTRACT'): assert ReviewReason.OCR_SOURCE in r.human_review.reasons
        save(out/(name+'_fixture_response.json'),{'synthetic':True,'llm_called':False,'response':fixture,'call_error':str(error) if error else None})
        save(out/(name+'_validated_result.json'),data)
        checks.append({'scenario':name,'status':'PASS','processing_status':data['processing_status'],
                       'review_required':r.human_review.required,'review_reasons':[r.value for r in r.human_review.reasons],
                       'evidence_status':to_evidence_status(r).value})
    assert before==fingerprint(),'code/KB changed during verification'
    report={'status':'PASS','checked_at':datetime.now().astimezone().isoformat(),'llm_calls':0,
            'scope':'Actual file parsers/chunker/local retrieval/prompt, fixed-response validator/review. No upload API, MySQL, UI or LLM server.',
            'kb_sha256':engine.kb_sha,'code_sha256':before,'file_cases':results,'preprocessing_rejected':rejected,
            'response_scenarios':checks,'utf8_cp949_scores_equal':True,
            'limits':['PNG/JPG use simulated OCR blocks; EasyOCR/model unavailable in inspected runtimes.',
                      'Synthetic selected candidates and confidence are protocol fixtures, not predicted labels or accuracy.',
                      'LLM transport, structured output enforcement, model quality, retries on a real endpoint and full upload/DB/UI remain unverified.']}
    save(out/'summary.json',report)
    lines=['# LLM 연결 전 통합 검증','',f"- 결과: PASS / {report['checked_at']}",
           '- 실제 파일 11건 중 정상 9건은 최신 파서→청킹→로컬 검색→판단 입력→프롬프트 생성 통과. 빈 파일·손상 파일 2건은 검색 전에 차단.',
           '- PNG/JPG는 합성 OCR 블록 2건으로 청킹 이후 연결만 검증. 실제 이미지 인식은 실행하지 않음.',
           '- 고정 응답 10종으로 실제 검증·검토 코드 실행. LLM 호출 0회. 업로드 API·MySQL·화면 연결은 포함하지 않음.',
           '- UTF-8/CP949 동일 본문의 순위·점수 일치, 긴 문서 분할, PDF 페이지/XLSX 시트/PPTX 슬라이드 1·2 보존 확인.',
           '', '| 사례 | 구분 | 청크 수 | 위치 | 출처 |', '|---|---|---:|---|---|']
    for r in results: lines.append(f"| {r['case_id']} | {r['origin']} | {r['chunks']} | {r['pages']} | {r['sources']} |")
    lines+=['','| 고정 응답 | 처리 상태 | 검토 필요 | 사유 |','|---|---|---|---|']
    for c in checks: lines.append(f"| {c['scenario']} | {c['processing_status']} | {c['review_required']} | {', '.join(c['review_reasons'])} |")
    lines+=['','프롬프트와 시험 응답에는 실제 모델 결과로 오인하지 않도록 표시했다. 후보 선택과 confidence는 시험 상수이며 검색·LLM의 정확도 지표로 사용하면 안 된다.',
            '', '## 실제 LLM 연결 시 필요한 것', '',
            '- 서버 종류·주소·모델 식별자 및 인증 설정',
            '- 요청/응답 구조, JSON Schema 강제 전달 방식, 토큰·생성 설정',
            '- 실제 정상·무관·복수·불확실 사례 실행과 원문 인용 검증',
            '- 타임아웃·오류·재시도와 업로드/DB/화면 흐름 확인','']
    (out/'summary.md').write_text('\n'.join(lines),encoding='utf-8')
    print(json.dumps({'status':'PASS','report':str(out/'summary.md'),'llm_calls':0},ensure_ascii=False))


if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--parser-python')
    p.add_argument('--worker',action='store_true',help=argparse.SUPPRESS)
    p.add_argument('--output-dir',type=Path,default=ROOT/'reports'/('pre_llm_'+datetime.now().strftime('%Y-%m-%d')))
    a=p.parse_args()
    out=a.output_dir.resolve()
    if a.worker: print(json.dumps(worker(out),ensure_ascii=False))
    elif not a.parser_python: p.error('--parser-python is required')
    else:
        try:
            main(a.parser_python,out)
        except Exception as exc:
            save(out/'summary.json',{'status':'FAIL','llm_calls':0,'error':str(exc)})
            (out/'summary.md').write_text('# 연결 전 검사 실패\n\n'+str(exc)+'\n',encoding='utf-8')
            raise
