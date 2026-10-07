"""Local HTTP health and parser/chunker smoke; no DB or LLM requests."""
import argparse
import importlib
from importlib import metadata
import json
from pathlib import Path
import platform
import sys

ROOT = Path(__file__).resolve().parents[2]


def run(work):
    results = []

    def check(name, fn):
        print('CHECK ' + name, flush=True)
        try:
            detail = fn()
            results.append({'name': name, 'passed': True, 'detail': detail})
        except Exception as exc:
            results.append({'name': name, 'passed': False,
                            'error_type': type(exc).__name__, 'message': str(exc)})
        print(('PASS ' if results[-1]['passed'] else 'FAIL ') + name, flush=True)

    for folder in ('phase1_검색', 'phase1_입력', 'phase1_입력/chunking', 'phase1_통합'):
        sys.path.insert(0, str(ROOT / folder))

    def web():
        from fastapi.testclient import TestClient
        api = importlib.import_module('api')
        assert Path(api.__file__).resolve() == ROOT / 'phase1_통합/api.py'
        with TestClient(api.app) as client:
            response = client.get('/health')
            assert response.status_code == 200 and response.json()['status'] == 'ok'
            readiness = client.get('/api/readiness')
            assert readiness.status_code == 200
            page = client.get('/')
            assert page.status_code == 200
        return {'health_status': 200, 'page_status': 200,
                'readiness_ready': readiness.json()['ready'],
                'not_ready': [r['name'] for r in readiness.json()['checks'] if not r['ok']]}

    check('web_health_page_and_readiness', web)
    for module in ('easyocr', 'torchvision', 'pdfplumber', 'docx', 'pptx', 'openpyxl'):
        check('import_' + module, lambda m=module: {'module': importlib.import_module(m).__name__})

    def text_case(name, content, expected_error=None):
        from parser_dispatcher import parse_file
        from chunker import make_chunks
        from chunk_format import validate_chunks
        path = work / name
        path.write_bytes(content)
        parsed = parse_file(path)
        if expected_error:
            assert expected_error in parsed['errors'], parsed['errors']
            return {'errors': parsed['errors']}
        assert not parsed['errors'], parsed['errors']
        chunks = make_chunks(parsed, 'E9999', 1)
        assert chunks
        validation = validate_chunks(chunks)
        assert not validation, validation
        return {'block_count': len(parsed['blocks']), 'chunk_count': len(chunks)}

    check('txt_parse_and_chunk', lambda: text_case('normal.txt', '퇴직자 계정은 삭제하고 처리 이력을 기록한다.'.encode('utf-8')))
    check('empty_txt', lambda: text_case('empty.txt', b'', 'empty_document'))
    check('corrupted_docx', lambda: text_case('broken.docx', b'not a zip document', 'corrupted_file'))
    check('corrupted_pdf', lambda: text_case('broken.pdf', b'%PDF-broken', 'corrupted_file'))

    def docx_case():
        from docx import Document
        from parser_dispatcher import parse_file
        from chunker import make_chunks
        from chunk_format import validate_chunks
        path = work / 'normal.docx'
        doc = Document()
        doc.add_paragraph('시험용 계정 관리 절차')
        doc.add_paragraph('계정은 부서장 승인 후 발급하고 퇴직 시 삭제한다.')
        doc.save(path)
        parsed = parse_file(path)
        assert not parsed['errors'], parsed['errors']
        chunks = make_chunks(parsed, 'E9999', 1)
        assert chunks and not validate_chunks(chunks)
        return {'block_count': len(parsed['blocks']), 'chunk_count': len(chunks)}

    check('docx_parse_and_chunk', docx_case)
    return {'scope': 'IN_PROCESS_HTTP_AND_PARSER_SMOKE', 'python': platform.python_version(),
            'results': results, 'passed': all(r['passed'] for r in results),
            'database_tested': False, 'llm_tested': False, 'ocr_inference_tested': False,
            'upload_to_report_e2e': False,
            'packages': sorted([{'name': d.metadata['Name'], 'version': d.version}
                                for d in metadata.distributions()], key=lambda d: d['name'].lower())}


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--work', type=Path, required=True, help='New scratch directory outside source')
    args = parser.parse_args()
    args.work.mkdir(parents=True, exist_ok=False)
    report = run(args.work)
    (args.work / 'report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps({k: v for k, v in report.items() if k != 'packages'}, ensure_ascii=False, indent=2))
    raise SystemExit(0 if report['passed'] else 1)

