"""
pptx_parser.py : PPTX 슬라이드 텍스트 추출 (Phase 1 입력 - 파워포인트 담당)

하는 일
  파워포인트 파일을 슬라이드 순서대로 읽어서 "슬라이드 제목(heading)", "글상자 내용(paragraph)",
  "표(table)"를 뽑고, 각각 몇 번째 슬라이드였는지 기록한 결과(딕셔너리)를 돌려줍니다.

팀 합의 규칙 (v2) - PDF/DOCX/XLSX 와 같은 형식
  1. block_type : paragraph, table, heading 세 가지만
  2. order      : 문서 내 등장 순서, 1부터 연속 (block_id 없음)
  3. page       : PPT는 "슬라이드 번호"(1부터). page_count 는 슬라이드 개수
  4. level      : heading 일 때만 1, 2, 3 중 하나, 나머지는 null
  5. table      : block_type 이 table 일 때만 {"rows": [...]}, 아니면 null
  6. table.rows : 2차원 배열, 병합 셀은 왼쪽 위 칸에만 값, 나머지는 빈 문자열 ""
  7. text       : paragraph/heading 일 때만 채우고, table 이면 ""
  8. 빈 블록은 버림 (빈 글상자, 모든 칸이 빈 표, 아무것도 없는 슬라이드)
  9. 손상/빈 문서 : 예외를 던지지 않고 blocks: [] + errors 에 사유 코드
 10. errors     : 항상 리스트, 없으면 []
 11. 일부(슬라이드 하나)만 실패하면 그 슬라이드만 건너뛰고 나머지를 뽑음 (errors 는 [])

이 파서가 정한 것 (★ 팀 확인 필요)
  ★ page          = 슬라이드 번호, page_count = 슬라이드 개수
  ★ 슬라이드 제목 = heading level 1 (제목 틀에 들어 있는 글)
  ★ 도형 순서     = 위에서 아래, 같은 높이면 왼쪽에서 오른쪽 (PPT 는 읽는 순서가 정해져 있지 않음)
  ★ 발표자 노트   = 슬라이드 맨 뒤에 paragraph 로 넣되 "[발표자 노트]" 를 앞에 붙임
  ★ 숨긴 슬라이드도 포함 (증적에서 숨긴 내용이 중요할 수 있어서)
  ★ 글상자 안의 문단은 각각 한 블록 (글머리표 한 줄 = 한 블록)
  ★ 한쇼 파일 대응: 열리긴 하는데 내용이 비어 보이면 AlternateContent 블록을 풀고 다시 읽음

errors 코드 (팀 합의: 두 가지만 사용)
  corrupted_file   파일을 열 수 없음
                   (깨진 파일, PPT가 아닌 파일, 0바이트, 비밀번호 걸린 파일, 없는 파일, 옛날 .ppt)
  empty_document   열리긴 하는데 뽑을 글자·표가 하나도 없음
                   (빈 슬라이드만 있는 문서, 캡처 그림만 붙여 넣은 문서)

사용법
  from pptx_parser import parse_pptx
  result = parse_pptx("C:/team/database/data/evidence/000003_v1.pptx")

  명령창에서 결과 보기:  python pptx_parser.py 파일경로.pptx
  (파일 이름에 띄어쓰기가 있으면 "따옴표"로 감싸세요)
"""

import io
import json
from collections import Counter
import os
import re
import sys
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

INCLUDE_NOTES = True                  # 발표자 노트를 결과에 넣을지
NOTES_PREFIX = "[발표자 노트] "        # 노트 앞에 붙이는 표시
SKIP_HIDDEN = False                   # True 로 바꾸면 숨긴 슬라이드를 건너뜀
TITLE_AS_HEADING = True               # 슬라이드 제목을 heading 으로 넣을지
HEADING_MAX_CHARS = 60                # 제목은 이 글자 수 이하

A_TEXT = "{http://schemas.openxmlformats.org/drawingml/2006/main}t"   # 도형 XML 안의 글자


def parse_pptx(path):
    """파워포인트 파일 하나를 읽어서 팀 합의 포맷의 딕셔너리로 돌려줍니다. 예외를 밖으로 던지지 않아요."""
    path = Path(path)
    result = {
        "source_file": path.name,
        "file_type": "pptx",
        "page_count": 0,
        "blocks": [],
        "errors": [],
    }

    # ① 파일이 없거나 0바이트면 열 수 없음
    if not path.exists() or os.path.getsize(path) == 0:
        result["errors"].append("corrupted_file")
        return result

    # ② 열기 (깨진 파일, 비밀번호 걸린 파일, 옛날 .ppt 는 여기서 실패)
    prs = _load(path)
    if prs is None:
        result["errors"].append("corrupted_file")
        return result

    result["page_count"] = len(prs.slides)
    blocks = _collect(prs)

    # 한쇼(한글과컴퓨터) 파일은 도형이 AlternateContent 블록 안에 들어 있어서
    # 파일은 열리는데 내용이 통째로 비어 보일 수 있어요. 그때 한 번 더 시도합니다.
    if not blocks:
        try:
            prs2 = Presentation(_unwrap_alternate_content(path))
            blocks = _collect(prs2)
            if blocks:
                result["page_count"] = len(prs2.slides)
        except Exception:
            pass

    if not blocks:
        result["errors"].append("empty_document")
        return result

    result["blocks"] = blocks
    return result


# ─────────────────────────────────────────────
# 내부 도우미 함수
# ─────────────────────────────────────────────
def _collect(prs):
    """슬라이드를 돌면서 블록 목록을 만듭니다."""
    blocks = []
    for slide_no, slide in enumerate(prs.slides, start=1):
        try:
            if SKIP_HIDDEN and _is_hidden(slide):
                continue
            items = _slide_items(slide)
        except Exception:
            continue                                  # 이 슬라이드만 건너뜀 (규칙 11)
        for kind, payload, level in items:
            if kind == "table":
                blocks.append(_block(len(blocks) + 1, "table", slide_no, table={"rows": payload}))
            elif kind == "heading":
                blocks.append(_block(len(blocks) + 1, "heading", slide_no, text=payload, level=level))
            else:
                blocks.append(_block(len(blocks) + 1, "paragraph", slide_no, text=payload))
    return blocks


ALTERNATE = re.compile(rb"<(?:\w+:)?AlternateContent\b.*?</(?:\w+:)?AlternateContent>", re.S)
FALLBACK = re.compile(rb"<(?:\w+:)?Fallback\b[^>]*>(.*?)</(?:\w+:)?Fallback>", re.S)
CHOICE = re.compile(rb"<(?:\w+:)?Choice\b[^>]*>(.*?)</(?:\w+:)?Choice>", re.S)


def _load(path):
    """
    파일을 엽니다. 못 열면 None.
    한쇼(한글과컴퓨터)에서 만든 파일은 AlternateContent 라는 특수 블록 때문에 열다가 실패할 수 있어서,
    실패하면 그 블록을 풀어낸 사본을 메모리에 만들어 다시 열어봅니다. (엑셀 파서와 같은 방식)
    """
    try:
        return Presentation(str(path))
    except Exception:
        pass
    try:
        return Presentation(_unwrap_alternate_content(path))
    except Exception:
        return None


def _unwrap_alternate_content(path):
    """AlternateContent 블록을 그 안의 Fallback(없으면 Choice) 내용으로 바꾼 사본을 메모리에 만듭니다."""
    def unwrap(match):
        inner = FALLBACK.search(match.group(0)) or CHOICE.search(match.group(0))
        return inner.group(1) if inner else b""

    src = zipfile.ZipFile(str(path))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename.endswith(".xml"):
                data = ALTERNATE.sub(unwrap, data)
            out.writestr(item, data)
    buf.seek(0)
    return buf


def _block(order, block_type, page, text="", level=None, table=None):
    return {"order": order, "block_type": block_type, "level": level,
            "page": page, "text": text, "table": table}


def _is_hidden(slide):
    return slide.element.get("show") == "0"


def _slide_items(slide):
    """
    슬라이드 하나에서 (종류, 내용, level) 목록을 위에서 아래 순서로 돌려줍니다.
    PPT 는 도형에 "읽는 순서"가 없어서, 도형 위치(위→아래, 같으면 왼→오른쪽)로 정렬해요.
    그룹으로 묶인 도형은 좌표가 "그룹 안 좌표"라서, 슬라이드 좌표로 바꿔서 비교합니다.
    """
    title_shape = slide.shapes.title if TITLE_AS_HEADING else None
    items = []                                   # (세로, 가로, 종류, 내용, level, 글자크기)

    for shape, (top, left) in _walk(slide.shapes):

        if getattr(shape, "has_chart", False):       # 차트 → 항목·값을 표로
            rows = _chart_rows(shape)
            if rows:
                items.append((top, left, "table", rows, None, None))
            continue

        if shape.has_table:
            rows = _table_rows(shape.table)
            if rows:                             # 모든 칸이 빈 표는 버림 (규칙 8)
                items.append((top, left, "table", rows, None, None))
            continue

        if shape.has_text_frame:
            is_title = title_shape is not None and shape == title_shape
            for text, size in _frame_texts(shape.text_frame):
                items.append((top, left, "heading" if is_title else "text", text, 1 if is_title else None, size))
            continue

        for text in _xml_texts(shape) or _diagram_texts(shape, slide.part):   # 스마트아트 등
            items.append((top, left, "text", text, None, None))

    items.sort(key=lambda x: (x[0], x[1]))
    if title_shape is None:
        items = _mark_heading_by_size(items)     # 제목 틀이 없는 PPT → 글자 크기로 제목 찾기
    out = [(kind, payload, level) for _, _, kind, payload, level, _ in items]

    if INCLUDE_NOTES and slide.has_notes_slide:  # 발표자 노트는 맨 뒤에
        frame = slide.notes_slide.notes_text_frame
        if frame is not None:
            out += [("text", NOTES_PREFIX + text, None) for text, _ in _frame_texts(frame)]
    return out


def _walk(shapes, offset=(0, 0), scale=(1.0, 1.0)):
    """
    그룹으로 묶인 도형 안까지 들어가서 (도형, 슬라이드 위 위치)를 하나씩 돌려줍니다.
    그룹 안 도형의 top/left 는 "그룹 안 좌표"라서, 그룹의 위치·크기로 환산해야 화면 순서가 맞아요.
    """
    for shape in shapes:
        top = (shape.top or 0) * scale[1] + offset[1]
        left = (shape.left or 0) * scale[0] + offset[0]
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(shape.shapes, *_group_transform(shape, offset, scale))
        else:
            yield shape, (top, left)


def _group_transform(group, offset, scale):
    """그룹 안 좌표 → 슬라이드 좌표 변환값(더할 값, 곱할 값)을 계산합니다."""
    try:
        xfrm = group.element.find(
            "{http://schemas.openxmlformats.org/drawingml/2006/main}xfrm"
        ) if group.element.find("{http://schemas.openxmlformats.org/drawingml/2006/main}xfrm") is not None else group.element.xfrm
        sx = xfrm.ext.cx / xfrm.chExt.cx if xfrm.chExt.cx else 1.0
        sy = xfrm.ext.cy / xfrm.chExt.cy if xfrm.chExt.cy else 1.0
        off_x = (group.left or 0) * scale[0] + offset[0] - xfrm.chOff.x * sx * scale[0]
        off_y = (group.top or 0) * scale[1] + offset[1] - xfrm.chOff.y * sy * scale[1]
        return (off_x, off_y), (sx * scale[0], sy * scale[1])
    except Exception:
        return offset, scale


def _mark_heading_by_size(items):
    """
    제목 틀이 없는 PPT(템플릿으로 만든 자료)에서 제목을 찾습니다.
      - 후보 : 글자 크기를 아는 글 중 2~60자이고, 숫자·기호만 있는 게 아닌 것
               (01, 02 같은 장식 번호는 제외)
      - 그중 가장 큰 글자 크기를 고르고, 그 크기인 글 중 가장 위에 있는 것 하나만 heading level 1
    본문 글자에 크기 지정이 없는 PPT 가 많아서, 본문과 비교하지 않고 "가장 큰 글자"로 판단해요.
    """
    candidates = [(top, i, size) for i, (top, left, kind, payload, level, size) in enumerate(items)
                  if kind == "text" and size and 2 <= len(payload) <= HEADING_MAX_CHARS
                  and not re.fullmatch(r"[\d\W_]+", payload)]
    if not candidates:
        return items
    biggest = max(size for _, _, size in candidates)
    top, i, _ = min((c for c in candidates if c[2] == biggest), key=lambda c: c[0])
    t, l, _, payload, _, size = items[i]
    items[i] = (t, l, "heading", payload, 1, size)
    return items


def _frame_texts(text_frame):
    """글상자 안의 문단을 (글, 글자크기) 로 한 줄씩 돌려줍니다. 빈 문단은 버려요 (규칙 8)."""
    out = []
    for paragraph in text_frame.paragraphs:
        text = _clean("".join(run.text for run in paragraph.runs) or paragraph.text)
        if not text:
            continue
        sizes = [run.font.size.pt for run in paragraph.runs if run.font.size]
        out.append((text, max(sizes) if sizes else None))
    return out


def _chart_rows(shape):
    """
    차트에서 항목 이름과 값을 표로 만듭니다.
      [["구분", "점검결과"], ["적합", "8"], ["미흡", "2"]]
    그림으로만 보이던 숫자를 글자로 남겨서 검색·분석에 쓸 수 있게 해요.
    """
    try:
        chart = shape.chart
        series = list(chart.plots[0].series)
        categories = [_clean(c) for c in chart.plots[0].categories]
        header = ["구분"] + [_clean(sr.name or f"계열{i + 1}") for i, sr in enumerate(series)]
        rows = [header]
        for r, category in enumerate(categories):
            values = [_number(sr.values[r]) for sr in series]
            if category or any(values):
                rows.append([category] + values)
        return rows if len(rows) > 1 else []
    except Exception:
        return []


def _number(value):
    """차트 값 → 글자 (정수는 소수점 없이)"""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


DIAGRAM_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/diagramData"


def _diagram_texts(shape, part):
    """스마트아트(도해)의 글자를 꺼냅니다. 글자는 슬라이드가 아니라 별도 파일에 들어 있어요."""
    try:
        for rel in part.rels.values():
            if rel.reltype == DIAGRAM_REL and rel.rId in shape.element.xml:
                return _texts_from_xml(rel.target_part.blob)
    except Exception:
        pass
    return []


def _texts_from_xml(blob):
    """XML 덩어리에서 글자(<a:t>)만 순서대로 뽑습니다."""
    from xml.etree import ElementTree
    try:
        root = ElementTree.fromstring(blob)
    except Exception:
        return []
    return [t for t in (_clean(node.text or "") for node in root.iter(A_TEXT)) if t]


def _xml_texts(shape):
    """python-pptx 가 글자로 안 주는 도형(스마트아트 등)에서 글자만 건져냅니다."""
    try:
        texts = [_clean(node.text or "") for node in shape.element.iter(A_TEXT)]
    except Exception:
        return []
    return [t for t in texts if t]


def _table_rows(table):
    """병합 칸은 왼쪽 위에만 값, 나머지는 "" (규칙 6). 완전히 빈 줄은 버림 (규칙 8)."""
    rows = []
    for row in table.rows:
        values = []
        for cell in row.cells:
            merged = getattr(cell, "is_spanned", False)      # 병합 칸 중 왼쪽 위가 아닌 칸
            values.append("" if merged else _clean(cell.text))
        if any(values):
            rows.append(values)
    return rows


def _clean(text):
    """줄바꿈·연속 공백을 하나의 공백으로 정리합니다."""
    return " ".join(str(text).split())


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("사용법: python pptx_parser.py 파일경로.pptx")
        raise SystemExit(1)
    print(json.dumps(parse_pptx(sys.argv[1]), ensure_ascii=False, indent=2))
