"""DOCX 추출 검사: 정상 문서(문단·표)와 손상/빈 DOCX (예외 대신 errors 로 반환)

실행:  입력 폴더에서  python -m pytest tests -v   (또는 python -m unittest discover -s tests)
"""
import io
from pathlib import Path
import shutil
import sys
import tempfile
import unittest
import zipfile

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx_extractor import extract_docx  # noqa: E402


class DocxTestCase(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.dir, ignore_errors=True)

    def path(self, name="t.docx"):
        return self.dir / name

    def save(self, doc, name="t.docx"):
        p = self.path(name)
        doc.save(str(p))
        return p

    def write_bytes(self, data, name="t.docx"):
        p = self.path(name)
        p.write_bytes(data)
        return p

    def assertErrorCode(self, path, code):
        """예외 없이 blocks 는 비고 errors 에 사유 코드가 담기는지 확인"""
        r = extract_docx(path)
        self.assertEqual(r["blocks"], [])
        self.assertIsInstance(r["errors"], list)
        self.assertEqual([e["code"] for e in r["errors"]], [code])
        self.assertTrue(r["errors"][0]["message"])
        self.assertEqual((r["file_type"], r["page_count"]), ("docx", None))


class ExtractTests(DocxTestCase):
    def test_top_level_format(self):
        doc = Document()
        doc.add_paragraph("본문")

        r = extract_docx(self.save(doc, "policy.docx"))

        self.assertEqual(
            r, {
                "source_file": "policy.docx", "file_type": "docx", "page_count": None,
                "blocks": [{"order": 1, "block_type": "paragraph", "level": None,
                            "page": None, "text": "본문", "table": None}],
                "errors": [],
            })

    def test_source_file_uses_original_name(self):
        doc = Document()
        doc.add_paragraph("본문")
        r = extract_docx(self.save(doc, "000001_v1.docx"), source_file="정책.docx")
        self.assertEqual(r["source_file"], "정책.docx")

    def test_paragraphs_and_headings(self):
        doc = Document()
        doc.add_heading("제1장 총칙", level=1)
        doc.add_paragraph("퇴직자 계정은 당일 삭제한다.")
        doc.add_paragraph("   ")  # 공백뿐인 문단은 뺀다
        doc.add_heading("1.1 계정", level=2)
        doc.add_heading("1.1.1 삭제", level=3)
        doc.add_heading("1.1.1.1 세부", level=4)  # 4단계 이하는 3으로 묶는다

        blocks = extract_docx(self.save(doc))["blocks"]

        self.assertEqual([(b["order"], b["block_type"], b["level"], b["text"]) for b in blocks], [
            (1, "heading", 1, "제1장 총칙"),
            (2, "paragraph", None, "퇴직자 계정은 당일 삭제한다."),
            (3, "heading", 2, "1.1 계정"),
            (4, "heading", 3, "1.1.1 삭제"),
            (5, "heading", 3, "1.1.1.1 세부"),
        ])

    def test_field_rules_per_block_type(self):
        doc = Document()
        doc.add_heading("제목", level=1)
        doc.add_paragraph("본문")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "표"

        blocks = extract_docx(self.save(doc))["blocks"]

        for b in blocks:
            self.assertIn(b["block_type"], {"paragraph", "table", "heading"})
            self.assertIsNone(b["page"])
            self.assertEqual(b["level"] is not None, b["block_type"] == "heading")
            self.assertEqual(b["table"] is not None, b["block_type"] == "table")
            self.assertEqual(b["text"] == "", b["block_type"] == "table")

    def test_table_block(self):
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        for i, row in enumerate([["구분", "내용"], ["접근통제", "계정 삭제"]]):
            for j, text in enumerate(row):
                table.cell(i, j).text = text

        (block,) = extract_docx(self.save(doc))["blocks"]

        self.assertEqual(block["block_type"], "table")
        self.assertEqual(block["text"], "")
        self.assertIsNone(block["level"])
        self.assertEqual(block["table"], {"rows": [["구분", "내용"], ["접근통제", "계정 삭제"]]})

    def test_order_keeps_document_flow(self):
        doc = Document()
        doc.add_paragraph("표 앞 문단")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "표 내용"
        doc.add_paragraph("표 뒤 문단")

        blocks = extract_docx(self.save(doc))["blocks"]

        self.assertEqual([b["order"] for b in blocks], [1, 2, 3])
        self.assertEqual([b["block_type"] for b in blocks], ["paragraph", "table", "paragraph"])

    def test_merged_cells_become_empty_strings(self):
        doc = Document()
        table = doc.add_table(rows=3, cols=3)
        table.cell(0, 0).merge(table.cell(0, 2)).text = "제목"      # 가로 병합
        table.cell(1, 0).merge(table.cell(2, 0)).text = "세로"      # 세로 병합
        for i, j, text in [(1, 1, "a"), (1, 2, "b"), (2, 1, "c"), (2, 2, "d")]:
            table.cell(i, j).text = text

        rows = extract_docx(self.save(doc))["blocks"][0]["table"]["rows"]

        self.assertEqual(rows, [["제목", "", ""], ["세로", "a", "b"], ["", "c", "d"]])

    def test_multiline_cell_and_nested_table(self):
        doc = Document()
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        cell.text = "첫 줄"
        cell.add_paragraph("둘째 줄")
        inner = cell.add_table(rows=1, cols=2)
        inner.cell(0, 0).text = "안쪽1"
        inner.cell(0, 1).text = "안쪽2"

        rows = extract_docx(self.save(doc))["blocks"][0]["table"]["rows"]

        self.assertEqual(rows[0][0], "첫 줄\n둘째 줄\n안쪽1 | 안쪽2")

    def test_table_only_document_is_not_empty(self):
        doc = Document()
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "표만 있음"

        blocks = extract_docx(self.save(doc))["blocks"]

        self.assertEqual([b["block_type"] for b in blocks], ["table"])


class ErrorsFieldTests(DocxTestCase):
    def test_broken_file_keeps_source_file(self):
        r = extract_docx(self.write_bytes(b"not a docx"), source_file="깨진.docx")
        self.assertEqual(r["source_file"], "깨진.docx")

    def test_good_file_errors_is_empty_list(self):
        doc = Document()
        doc.add_paragraph("본문")
        r = extract_docx(self.save(doc))
        self.assertEqual(r["errors"], [])  # None 이 아니라 []
        self.assertEqual(len(r["blocks"]), 1)


class BrokenDocxTests(DocxTestCase):
    def test_missing_file(self):
        self.assertErrorCode(self.path("없는파일.docx"), "FILE_NOT_FOUND")

    def test_zero_byte_file(self):
        self.assertErrorCode(self.write_bytes(b""), "EMPTY_FILE")

    def test_random_bytes(self):
        self.assertErrorCode(self.write_bytes(b"\x00\x01\x02 not a docx " * 50), "FILE_PARSE_FAILED")

    def test_text_file_renamed_to_docx(self):
        self.assertErrorCode(self.write_bytes("그냥 텍스트 파일입니다".encode("utf-8")), "FILE_PARSE_FAILED")

    def test_truncated_docx(self):
        doc = Document()
        doc.add_paragraph("정상 문서")
        good = self.save(doc, "good.docx").read_bytes()
        self.assertErrorCode(self.write_bytes(good[: len(good) // 2]), "FILE_PARSE_FAILED")

    def test_zip_without_document_xml(self):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as z:
            z.writestr("hello.txt", "docx 가 아닌 zip")
        self.assertErrorCode(self.write_bytes(buf.getvalue()), "FILE_PARSE_FAILED")

    def test_corrupted_document_xml(self):
        doc = Document()
        doc.add_paragraph("정상 문서")
        good = self.save(doc, "good.docx")

        buf = io.BytesIO()
        with zipfile.ZipFile(good) as src, zipfile.ZipFile(buf, "w") as dst:
            for item in src.infolist():
                data = src.read(item.filename)
                if item.filename == "word/document.xml":
                    data = data[: len(data) // 2]  # XML 이 중간에서 끊김
                dst.writestr(item, data)

        self.assertErrorCode(self.write_bytes(buf.getvalue()), "FILE_PARSE_FAILED")


class EmptyDocxTests(DocxTestCase):
    def test_empty_table_block_is_dropped(self):
        doc = Document()
        doc.add_paragraph("본문")
        doc.add_table(rows=2, cols=2)  # 글자 없는 표
        blocks = extract_docx(self.save(doc))["blocks"]
        self.assertEqual([b["block_type"] for b in blocks], ["paragraph"])

    def test_document_with_no_content(self):
        self.assertErrorCode(self.save(Document()), "EMPTY_DOCUMENT")

    def test_only_blank_paragraphs(self):
        doc = Document()
        doc.add_paragraph("")
        doc.add_paragraph("   ")
        self.assertErrorCode(self.save(doc), "EMPTY_DOCUMENT")

    def test_only_empty_table(self):
        doc = Document()
        doc.add_table(rows=2, cols=2)
        self.assertErrorCode(self.save(doc), "EMPTY_DOCUMENT")


if __name__ == "__main__":
    unittest.main()
