"""
xlsx_parser.py : XLSX 시트·셀 텍스트 추출 (Phase 1 입력 - 엑셀 담당)

하는 일
  엑셀 파일을 시트 순서대로 읽어서 "시트 이름(heading)"과 "표(table)"를 뽑고,
  각각 몇 번째 시트에 있었는지 기록한 결과(딕셔너리)를 돌려줍니다.

팀 합의 규칙 (v2) - PDF/DOCX 와 같은 형식
  1. block_type : paragraph, table, heading 세 가지만
  2. order      : 문서 내 등장 순서, 1부터 연속 (block_id 없음)
  3. page       : 엑셀은 "시트 번호"(1부터). page_count 는 시트 개수
  4. level      : heading 일 때만 1, 2, 3 중 하나, 나머지는 null
  5. table      : block_type 이 table 일 때만 {"rows": [...]}, 아니면 null
  6. table.rows : 2차원 배열, 병합 셀은 왼쪽 위 칸에만 값, 나머지는 빈 문자열 ""
  7. text       : paragraph/heading 일 때만 채우고, table 이면 ""
  8. 빈 블록은 버림 (빈 줄, 빈 칸(열), 내용이 하나도 없는 시트)
  9. 손상/빈 문서 : 예외를 던지지 않고 blocks: [] + errors 에 사유 코드
 10. errors     : 항상 리스트, 없으면 []
 11. 일부(시트 하나)만 실패하면 그 시트만 건너뛰고 나머지를 뽑음 (errors 는 [])

이 파서가 정한 것 (★ 팀 확인 필요)
  ★ page       = 시트 번호, page_count = 시트 개수 (DOCX 는 null, PDF 는 실제 페이지)
  ★ 시트 이름  = heading level 1 블록으로 넣음 (형식에 시트 이름 칸이 없어서)
  ★ 숨긴 시트·숨긴 행도 포함 (증적에서 숨긴 내용이 중요할 수 있어서)
  ★ 수식은 계산된 값으로. 저장된 계산값이 없으면 수식 글자(=SUM(...))를 넣음
  ★ 아주 큰 시트는 MAX_TABLE_ROWS 줄씩 나누고, 나뉜 표마다 머리글 줄을 다시 붙임
  ★ 표 맨 윗줄이 "한 칸만 채워진 제목 줄"이면 표에서 빼서 heading level 2 블록으로 만듦
    (엑셀 증적은 표 위에 병합된 제목 줄이 흔한데, 그대로 두면 청킹이 제목 줄을 머리글로 착각함)

errors 코드 (팀 합의: 두 가지만 사용)
  corrupted_file   파일을 열 수 없음
                   (깨진 파일, 엑셀이 아닌 파일, 0바이트, 비밀번호 걸린 파일, 없는 파일, 옛날 .xls)
  empty_document   열리긴 하는데 뽑을 글자가 하나도 없음 (빈 시트만 있는 문서)

사용법
  from xlsx_parser import parse_xlsx
  result = parse_xlsx("C:/team/database/data/evidence/000002_v1.xlsx")

  명령창에서 결과 보기:  python xlsx_parser.py 파일경로.xlsx
  (파일 이름에 띄어쓰기가 있으면 "따옴표"로 감싸세요)
"""

import datetime
import io
import json
import os
import re
import sys
import zipfile
from pathlib import Path

import openpyxl

MAX_TABLE_ROWS = 200        # 표 한 블록의 최대 줄 수 (이보다 길면 나눔)
REPEAT_HEADER = True        # 나눈 표마다 첫 줄(머리글)을 다시 붙일지
SKIP_HIDDEN = False         # True 로 바꾸면 숨긴 시트·숨긴 행을 건너뜀
SHEET_NAME_AS_HEADING = True   # 시트 이름을 heading 블록으로 넣을지


def parse_xlsx(path):
    """엑셀 파일 하나를 읽어서 팀 합의 포맷의 딕셔너리로 돌려줍니다. 예외를 밖으로 던지지 않아요."""
    path = Path(path)
    result = {
        "source_file": path.name,
        "file_type": "xlsx",
        "page_count": 0,
        "blocks": [],
        "errors": [],
    }

    # ① 파일이 없거나 0바이트면 열 수 없음
    if not path.exists() or os.path.getsize(path) == 0:
        result["errors"].append("corrupted_file")
        return result

    # ② 열기 (깨진 파일, 비밀번호 걸린 파일, 엑셀이 아닌 파일, 옛날 .xls 는 여기서 실패)
    #    data_only=True  : 수식의 "계산된 값"
    #    data_only=False : 수식 글자 그대로 (계산값이 저장 안 된 파일 대비용)
    wb_value = _load(path, data_only=True)
    if wb_value is None:
        result["errors"].append("corrupted_file")
        return result
    wb_formula = _load(path, data_only=False)

    blocks = []
    sheets = [ws for ws in wb_value.worksheets]
    result["page_count"] = len(sheets)

    for sheet_no, ws in enumerate(sheets, start=1):
        try:
            if SKIP_HIDDEN and ws.sheet_state != "visible":
                continue
            ws_formula = wb_formula[ws.title] if wb_formula is not None and ws.title in wb_formula.sheetnames else None
            rows = _sheet_rows(ws, ws_formula)
        except Exception:
            continue                      # 이 시트만 건너뜀 (규칙 11)
        if not rows:                      # 빈 시트는 버림 (규칙 8)
            continue

        if SHEET_NAME_AS_HEADING and ws.title.strip():
            blocks.append(_block(len(blocks) + 1, "heading", sheet_no, text=ws.title.strip(), level=1))
        rows, titles = _peel_title_rows(rows)
        for title in titles:                       # 표 위의 제목 줄 → heading
            blocks.append(_block(len(blocks) + 1, "heading", sheet_no, text=title, level=2))
        for part in _split_rows(rows):
            blocks.append(_block(len(blocks) + 1, "table", sheet_no, table={"rows": part}))

    if not blocks:
        result["errors"].append("empty_document")
        return result

    result["blocks"] = blocks
    return result


# ─────────────────────────────────────────────
# 내부 도우미 함수
# ─────────────────────────────────────────────
ALTERNATE = re.compile(rb"<(?:\w+:)?AlternateContent\b.*?</(?:\w+:)?AlternateContent>", re.S)
FALLBACK = re.compile(rb"<(?:\w+:)?Fallback\b[^>]*>(.*?)</(?:\w+:)?Fallback>", re.S)
CHOICE = re.compile(rb"<(?:\w+:)?Choice\b[^>]*>(.*?)</(?:\w+:)?Choice>", re.S)


def _load(path, data_only):
    """
    엑셀 파일을 엽니다. 못 열면 None.
    한셀(한글과컴퓨터)에서 만든 엑셀은 스타일 정보 안에 AlternateContent 라는 특수 블록을 넣는데,
    openpyxl 이 이 블록을 건너뛰어서 스타일 개수가 모자라 열다가 오류가 납니다.
    그래서 처음 열기에 실패하면, 이 블록을 풀어낸 사본을 메모리에 만들어서 다시 열어봅니다.
    """
    try:
        return openpyxl.load_workbook(str(path), data_only=data_only)
    except Exception:
        pass
    try:
        return openpyxl.load_workbook(_unwrap_alternate_content(path), data_only=data_only)
    except Exception:
        return None


def _unwrap_alternate_content(path):
    """AlternateContent 블록을 그 안의 Fallback(없으면 Choice) 내용으로 바꾼 사본을 메모리에 만듭니다."""
    def unwrap(match):
        inner = FALLBACK.search(match.group(0)) or CHOICE.search(match.group(0))
        return inner.group(1) if inner else b""

    src = zipfile.ZipFile(str(path))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename.endswith(".xml"):
                data = ALTERNATE.sub(unwrap, data)
            out.writestr(item, data)
    buf.seek(0)
    return buf


def _block(order, block_type, page, text="", level=None, table=None):
    return {"order": order, "block_type": block_type, "level": level,
            "page": page, "text": text, "table": table}


def _sheet_rows(ws, ws_formula):
    """
    시트 하나를 2차원 배열로 만듭니다.
      - 병합 셀은 openpyxl 이 왼쪽 위 칸에만 값을 주고 나머지는 None → "" 이 됩니다 (규칙 6)
      - 수식은 계산된 값, 계산값이 없으면 수식 글자
      - 완전히 빈 줄과 완전히 빈 칸(열)은 버립니다 (규칙 8)
    """
    grid = []
    for r, row in enumerate(ws.iter_rows(), start=1):
        if SKIP_HIDDEN and ws.row_dimensions[r].hidden:
            continue
        values = []
        for cell in row:
            text = _cell_text(cell.value)
            if not text and ws_formula is not None:            # 계산값이 없으면 수식 글자로
                text = _cell_text(ws_formula.cell(row=cell.row, column=cell.column).value)
            values.append(text)
        grid.append(values)

    grid = [row for row in grid if any(row)]                   # 빈 줄 버림
    if not grid:
        return []
    width = max(len(row) for row in grid)
    grid = [row + [""] * (width - len(row)) for row in grid]   # 줄 길이 맞추기
    keep = [k for k in range(width) if any(row[k] for row in grid)]   # 빈 칸(열) 버림
    return [[row[k] for k in keep] for row in grid]


def _peel_title_rows(rows):
    """
    표 맨 위의 "한 칸만 채워진 줄"을 제목으로 빼냅니다. 돌려주는 값: (남은 rows, 제목 목록)
      예) 1줄 ["[채용정보] 7월 2주차...", "", "", ...]   ← 제목 (병합된 칸)
          2줄 ["연번", "회사명", "구인분야", ...]        ← 진짜 머리글
    이렇게 빼두면 청킹에서 "첫 줄 = 머리글"로 짝지어도 맞아요.
    칸이 1개뿐인 표는 건드리지 않아요.
    """
    titles = []
    while len(rows) > 1 and len(rows[0]) > 1 and sum(1 for c in rows[0] if c) == 1:
        titles.append(next(c for c in rows[0] if c))
        rows = rows[1:]
    return rows, titles


def _split_rows(rows):
    """아주 긴 표를 MAX_TABLE_ROWS 줄씩 나눕니다. 나뉜 표마다 머리글 줄을 다시 붙여요."""
    if len(rows) <= MAX_TABLE_ROWS:
        return [rows]
    header, body = rows[0], rows[1:]
    size = MAX_TABLE_ROWS - (1 if REPEAT_HEADER else 0)
    parts = [body[i:i + size] for i in range(0, len(body), size)]
    return [([header] + part if REPEAT_HEADER else part) for part in parts]


def _cell_text(value):
    """셀 값을 글자로 바꿉니다. 날짜는 보이는 대로, 정수는 소수점 없이, 빈 칸은 ""."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime.datetime):
        if (value.hour, value.minute, value.second) == (0, 0, 0):
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, datetime.date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, datetime.time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return " ".join(str(value).split())


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("사용법: python xlsx_parser.py 파일경로.xlsx")
        raise SystemExit(1)
    print(json.dumps(parse_xlsx(sys.argv[1]), ensure_ascii=False, indent=2))
